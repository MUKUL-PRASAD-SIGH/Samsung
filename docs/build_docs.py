"""Build the .docx files in docs/ from the Markdown sources in docs/src/.

    python docs/build_docs.py        # needs: pip install python-docx

Supports the restricted Markdown the sources use: headings, paragraphs, bullets, numbered lists, pipe tables, fenced code,
**bold** and `code`. Also writes docs/Kairos_Complete_Documentation.docx combining every chapter.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

HERE = Path(__file__).resolve().parent
SRC, OUT = HERE / "src", HERE
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def shade(cell, hex_fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    pr.append(el)


def runs(par, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 10.5) - 1)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            continue
        else:
            r = par.add_run(part); r.bold = bold
        if size:
            r.font.size = Pt(size)


def new_doc():
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for name, size in (("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12)):
        h = d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.color.rgb = ACCENT
    return d


def code_block(d, lines):
    t = d.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]; shade(c, "F2F4F7")
    c.paragraphs[0].text = ""
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def table(d, rows):
    n = max(len(r) for r in rows)
    t = d.add_table(rows=len(rows), cols=n); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(n):
            cell = t.rows[i].cells[j]; cell.paragraphs[0].text = ""
            runs(cell.paragraphs[0], row[j] if j < len(row) else "", size=9, bold=(i == 0))
            if i == 0:
                shade(cell, "DCE6F2")
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def render(d, text, first_heading_level=1):
    lines, i = text.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            buf, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            code_block(d, buf)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(d, rows); continue
        elif ln.startswith("#"):
            m = re.match(r"(#+)\s+(.*)", ln)
            d.add_heading(m.group(2), level=min(len(m.group(1)) - 1 + first_heading_level, 3))
        elif re.match(r"\s*[-*] ", ln):
            runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"\s*\d+\. ", ln):
            # Numbers are written out: Word's "List Number" style keeps counting across separate lists.
            num, rest = re.match(r"\s*(\d+)\. (.*)", ln).groups()
            par = d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            runs(par, f"{num}.  {rest}")
        elif ln.strip():
            runs(d.add_paragraph(), ln.strip())
        i += 1


def title_of(text):
    return re.match(r"#\s+(.*)", text).group(1)


def main():
    files = sorted(SRC.glob("*.md"))
    if not files:
        sys.exit("no sources in docs/src")
    combined = new_doc()
    combined.add_heading("Kairos · Complete Technical Documentation", 0)
    combined.add_paragraph("Samsung Hackathon, Theme 05: a full-duplex, interruptible real-time agent.")
    combined.add_paragraph("Contents").runs[0].bold = True
    for f in files:
        combined.add_paragraph(title_of(f.read_text(encoding="utf-8")), style="List Bullet")
    for f in files:
        text = f.read_text(encoding="utf-8")
        d = new_doc(); render(d, text)
        d.save(OUT / (f.stem + ".docx")); print("wrote", f.stem + ".docx")
        combined.add_page_break(); render(combined, text)
    combined.save(OUT / "Kairos_Complete_Documentation.docx"); print("wrote Kairos_Complete_Documentation.docx")


if __name__ == "__main__":
    main()
